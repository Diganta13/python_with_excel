# import Workbook from openpyxl
from openpyxl import Workbook

# import DoughnutChart, Reference from openpyxl.chart sub_module .
from openpyxl.chart import DoughnutChart, Reference

# import DataPoint from openpyxl.chart.series class
from openpyxl.chart.series import DataPoint

# Call a Workbook() function of openpyxl
# to create a new blank Workbook object
wb = Workbook()

# Get workbook active sheet
# from the active attribute.
ws = wb.active

# data given
data = [
	['Pie', 2014],
	['Plain', 40],
	['Jam', 2],
	['Lime', 20],
	['Chocolate', 30],
]

# write content of each row in 1st and 2nd
# column of the active sheet respectively .
for row in data:
	ws.append(row)

# Create object of DoughnutChart class
chart = DoughnutChart()

# create data for plotting
labels = Reference(ws, min_col = 1, min_row = 2, max_row = 5)
data = Reference(ws, min_col = 2, min_row = 1, max_row = 5)

# adding data to the Doughnut chart object
chart.add_data(data, titles_from_data = True)

# set labels in the chart object
chart.set_categories(labels)

# set the title of the chart
chart.title = "Doughnuts Chart"

# set style of the chart
chart.style = 26

# add chart to the sheet
# the top-left corner of a chart
# is anchored to cell E1 .
ws.add_chart(chart, "E1")

# save the file
wb.save("doughnut.xlsx")
